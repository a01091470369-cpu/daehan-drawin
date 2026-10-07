# -*- coding: utf-8 -*-
"""reading.json → 도면검토보고서 docx (범용. 법정 문서 아님)."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import setup
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x7A, 0x7A, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AL = {"L": WD_ALIGN_PARAGRAPH.LEFT, "C": WD_ALIGN_PARAGRAPH.CENTER}
MARK = {"확인": "●", "추정": "△", "미확인": "○"}


def _f(r, size, bold=False, color=None):
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    rpr = r._r.get_or_add_rPr(); rf = OxmlElement("w:rFonts")
    for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rf.set(qn(a), FONT)
    rpr.insert(0, rf)


def _p(d, runs, al="L", sb=2, sa=2, rule=None):
    p = d.add_paragraph(); p.alignment = AL[al]
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    for t, size, bold, color in (runs if isinstance(runs, list) else [(runs, 10.5, False, None)]):
        _f(p.add_run(t), size, bold, color)
    if rule:
        pPr = p._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
        bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), str(rule))
        bt.set(qn("w:space"), "6"); bt.set(qn("w:color"), "1F3864")
        pb.append(bt); pPr.append(pb)
    return p


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), fill); tcPr.append(sh)


def _cell(cell, txt, size=9, bold=False, al="C", color=None):
    cell.text = ""
    p = cell.paragraphs[0]; p.alignment = AL[al]
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    _f(p.add_run(str(txt)), size, bold, color)
    tcPr = cell._tc.get_or_add_tcPr(); v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center"); tcPr.append(v)


def _borders(t):
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement("w:" + e); x.set(qn("w:val"), "single")
        x.set(qn("w:sz"), "4"); x.set(qn("w:color"), "BFC7D2"); b.append(x)
    t._tbl.tblPr.append(b)


def _table(d, headers, rows, widths, al=None):
    t = d.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t)
    for ci, h in enumerate(headers):
        c = t.cell(0, ci); _cell(c, h, 9, True, "C", WHITE)
        _shade(c, "1F3864"); c.width = Inches(widths[ci])
    for ri, row in enumerate(rows, 1):
        for ci, v in enumerate(row):
            c = t.cell(ri, ci)
            _cell(c, v if v not in (None, "") else "—", 9, False, (al or ["C"] * len(headers))[ci])
            c.width = Inches(widths[ci])
    return t


def build(work=None):
    reading, work, out = setup(work)
    b = reading["건물"]; name = b["명칭"]
    d = Document()
    s = d.sections[0]
    s.page_width = Inches(8.27); s.page_height = Inches(11.69)
    s.left_margin = s.right_margin = Inches(0.79); s.top_margin = s.bottom_margin = Inches(0.71)

    _p(d, [("도 면 검 토 보 고 서", 24, True, NAVY)], al="C", sa=4)
    _p(d, [(name, 13, False, None)], al="C", sa=2)
    _p(d, [("※ 본 보고서는 도면 판독에 따른 내부 검토자료이며, 법정 점검보고서가 아닙니다.", 8.5, False, GRAY)],
       al="C", sa=12, rule=10)

    _p(d, [("Ⅰ.  건물 개요", 13, True, NAVY)], al="L", sb=6, sa=5)
    _table(d, ["구분", "내용"], [[k, v] for k, v in b.items() if k != "명칭"],
           [1.3, 5.3], al=["C", "L"])

    _p(d, [("Ⅱ.  검토 도면", 13, True, NAVY)], al="L", sb=12, sa=5)
    _table(d, ["id", "분야", "도면명 / 파일", "쪽"],
           [[x["id"], x["분야"], (x.get("도면명") or x["파일"]), x.get("페이지수", "")]
            for x in reading["도면"]],
           [0.5, 0.9, 4.5, 0.6], al=["C", "C", "L", "C"])

    if reading["계통"]:
        _p(d, [("Ⅲ.  계통 구성", 13, True, NAVY)], al="L", sb=12, sa=5)
        for c in reading["계통"]:
            _p(d, [(f"▪ [{c.get('분야','')}] {c.get('계통명','')}", 10.5, True, None)], al="L", sb=4, sa=1)
            if c.get("흐름"):
                _p(d, [("    " + c["흐름"], 10, False, NAVY)], al="L", sb=0, sa=1)
            if c.get("설명"):
                _p(d, [("    " + c["설명"], 9.5, False, None)], al="L", sb=0, sa=2)

    _p(d, [("Ⅳ.  설비 현황", 13, True, NAVY)], al="L", sb=12, sa=3)
    _p(d, [("확신도    ● 확인(도면에 명시)    △ 추정(교차검증 필요)    ○ 미확인", 8.5, False, GRAY)],
       al="L", sb=0, sa=4)
    by = {}
    for it in reading["설비"]:
        by.setdefault(it.get("분야", "기타"), []).append(it)
    for fld, items in by.items():
        _p(d, [(f"{fld}  ({len(items)}건)", 11, True, None)], al="L", sb=8, sa=3)
        _table(d, ["분류", "명칭", "규격", "수량", "단위", "위치", "출처", ""],
               [[i.get("분류", ""), i.get("명칭", ""), i.get("규격", ""), i.get("수량", ""),
                 i.get("단위", ""), i.get("위치", ""), i.get("출처", ""),
                 MARK.get(i.get("확신도", ""), "")] for i in items],
               [0.9, 1.5, 1.5, 0.45, 0.4, 1.1, 0.65, 0.25],
               al=["C", "L", "L", "C", "C", "L", "C", "C"])

    if reading["미확인"]:
        _p(d, [("Ⅴ.  미확인 항목 (추가 자료 필요)", 13, True, NAVY)], al="L", sb=12, sa=5)
        _table(d, ["항목", "사유", "필요자료"],
               [[u.get("항목", ""), u.get("사유", ""), u.get("필요자료", "")] for u in reading["미확인"]],
               [1.9, 2.5, 2.2], al=["L", "L", "L"])

    _p(d, [("도면분석기  ｜  판독 내용은 도면 기준이며, 수량·규격은 장비일람표와 대조 후 확정하십시오.",
            8.5, False, GRAY)], al="C", sb=14, rule=6)

    p = os.path.join(out, f"{name}_도면검토보고서.docx")
    d.save(p)
    print("생성:", p)
    return p


if __name__ == "__main__":
    build()
